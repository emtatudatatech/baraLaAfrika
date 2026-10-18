"""`python -m pipelines gold-price`: the monthly gold price into db/seeds/gold_price.json.

Primary: the World Bank's Commodity Price Data (the Pink Sheet), monthly averages in nominal
US dollars per troy ounce, CC BY 4.0. The file's link changes with each release, so it is read
from the commodity-markets page. Fallback: the IMF's Primary Commodity Price System (PCPS),
used only when the World Bank file cannot be had; it is also fetched as a cross-check, and months
where the two differ by more than the configured tolerance are reported. The two series are
never mixed: each month keeps its source. The price is a monthly average, not a spot price.
"""

from __future__ import annotations

import argparse
import datetime as dt
import io
import json
import re
from dataclasses import dataclass, field
from typing import Any

import openpyxl

from .gold_common import (
    GOLD_RAW,
    PRICE_SEED,
    GoldError,
    SourceRecord,
    config,
    generated,
    imf_month,
    imf_series,
    month_end,
    table,
    write_seed,
)
from .http import FetchError, fetch_bytes

WORLD_BANK = SourceRecord(
    "worldbank-pink-sheet",
    "World Bank Commodity Price Data (The Pink Sheet), monthly prices",
    "https://www.worldbank.org/en/research/commodity-markets",
)
IMF_PCPS = SourceRecord(
    "imf-pcps",
    "IMF Primary Commodity Price System (PCPS): gold, US dollars per troy ounce",
    "https://data.imf.org/",
)
MONTH_RE = re.compile(r"^(\d{4})M(\d{2})$")
UPDATED_RE = re.compile(r"Updated on ([A-Z][a-z]+ \d{1,2}, \d{4})")


@dataclass
class PinkSheet:
    updated: dt.date | None
    prices: dict[dt.date, float]
    description: str | None = None


@dataclass
class PriceResult:
    prices: dict[dt.date, float]
    source: SourceRecord
    as_of: dt.date
    warnings: list[str] = field(default_factory=list)
    pcps: dict[dt.date, float] = field(default_factory=dict)


def monthly_file_url(page_html: str, file_name: str) -> str:
    pattern = rf'https://thedocs\.worldbank\.org/[^"\'\s]*/{re.escape(file_name)}'
    links = sorted(set(re.findall(pattern, page_html)))
    if len(links) != 1:
        raise GoldError(f"Expected one link to {file_name} on the World Bank page, found {links}")
    return str(links[0])


def parse_pink_sheet(body: bytes, *, sheet: str, column: str, unit: str) -> PinkSheet:
    workbook = openpyxl.load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    if sheet not in workbook.sheetnames:
        raise GoldError(f"The Pink Sheet has no {sheet!r} sheet (found {workbook.sheetnames})")
    rows = list(workbook[sheet].iter_rows(values_only=True))
    updated = None
    header_index = None
    gold_col = None
    for index, row in enumerate(rows[:20]):
        for cell in row:
            if isinstance(cell, str) and (match := UPDATED_RE.search(cell)):
                updated = dt.datetime.strptime(match[1], "%B %d, %Y").date()
        if column in row:
            header_index, gold_col = index, row.index(column)
            break
    if header_index is None or gold_col is None:
        raise GoldError(f"No {column!r} column in the Pink Sheet's {sheet!r} sheet")
    found_unit = rows[header_index + 1][gold_col]
    if found_unit != unit:
        raise GoldError(f"The Pink Sheet's {column} unit is {found_unit!r}, expected {unit!r}")
    prices: dict[dt.date, float] = {}
    for row in rows[header_index + 2 :]:
        label = row[0]
        if not isinstance(label, str) or (match := MONTH_RE.match(label.strip())) is None:
            continue
        value = row[gold_col]
        if isinstance(value, int | float) and value > 0:
            prices[dt.date(int(match[1]), int(match[2]), 1)] = float(value)
    if not prices:
        raise GoldError("The Pink Sheet has no gold prices")
    return PinkSheet(updated, prices, _description(workbook, column))


def _description(workbook: Any, column: str) -> str | None:
    """The series definition from the 'Description' sheet, e.g. "Gold, spot average of ..."."""
    if "Description" not in workbook.sheetnames:
        return None
    for row in workbook["Description"].iter_rows(values_only=True):
        for cell in row:
            if isinstance(cell, str) and cell.startswith(f"{column}, "):
                return cell.strip()
    return None


def parse_pcps(body: bytes) -> dict[dt.date, float]:
    series = imf_series(json.loads(body))
    found = {key: values for key, values in series.items() if key[1:3] == ("PGOLD", "USD")}
    if len(found) != 1:
        raise GoldError(f"Expected one PGOLD USD series from the IMF, found {list(found)}")
    (values,) = found.values()
    return {imf_month(period): value for period, value in values.items()}


def cross_check(
    primary: dict[dt.date, float], other: dict[dt.date, float], tolerance_pct: float
) -> list[str]:
    warnings = []
    for month in sorted(set(primary) & set(other)):
        difference = abs(primary[month] - other[month]) / other[month] * 100
        if difference > tolerance_pct:
            warnings.append(
                f"{month:%Y-%m}: World Bank {primary[month]:,.2f} vs IMF {other[month]:,.2f} "
                f"({difference:.1f}% apart)"
            )
    return warnings


def gaps(prices: dict[dt.date, float]) -> list[dt.date]:
    """Months missing between the first and last price (a hole is reported, never filled)."""
    months = sorted(prices)
    missing = []
    current = months[0]
    while current < months[-1]:
        current = (current.replace(day=28) + dt.timedelta(days=4)).replace(day=1)
        if current not in prices:
            missing.append(current)
    return missing


def build(*, refresh: bool) -> PriceResult:
    settings = config()["price"]
    pcps: dict[dt.date, float] = {}
    pcps_error = None
    try:
        pcps_body = fetch_bytes(
            settings["pcps_url"], GOLD_RAW / "imf_pcps_gold.json", refresh=refresh, timeout=90
        )
        pcps = parse_pcps(pcps_body)
    except (FetchError, GoldError) as error:
        pcps_error = str(error)
    try:
        page = fetch_bytes(
            settings["page_url"], GOLD_RAW / "worldbank_commodity_markets.html", refresh=refresh
        ).decode("utf-8", errors="replace")
        url = monthly_file_url(page, settings["file_name"])
        body = fetch_bytes(url, GOLD_RAW / settings["file_name"], refresh=refresh, timeout=120)
        sheet = parse_pink_sheet(
            body, sheet=settings["sheet"], column=settings["column"], unit=settings["unit"]
        )
    except (FetchError, GoldError) as error:
        if not pcps:
            raise GoldError(
                f"No gold price: World Bank failed ({error}); IMF failed ({pcps_error})"
            ) from error
        return PriceResult(
            pcps,
            IMF_PCPS,
            month_end(max(pcps)),
            [f"World Bank Pink Sheet unavailable ({error}); the IMF PCPS series is used."],
            pcps,
        )
    warnings = cross_check(sheet.prices, pcps, settings["cross_check_tolerance_pct"])
    if pcps_error:
        warnings.append(f"IMF PCPS cross-check unavailable: {pcps_error}")
    if sheet.description:
        warnings.insert(0, f"Series definition (World Bank): {sheet.description}")
    return PriceResult(
        sheet.prices, WORLD_BANK, sheet.updated or month_end(max(sheet.prices)), warnings, pcps
    )


def dataset(result: PriceResult) -> dict[str, Any]:
    months = sorted(result.prices)
    if result.source is WORLD_BANK:
        name = "World Bank Commodity Price Data (The Pink Sheet): gold, monthly average"
        provider = "World Bank (Prospects Group)"
        license_text = "CC BY 4.0 (World Bank dataset terms)"
        attribution = (
            "The World Bank: Commodity Price Data (The Pink Sheet), monthly prices: "
            "World Bank Prospects Group. CC BY 4.0."
        )
    else:
        name = "IMF Primary Commodity Price System: gold, monthly average"
        provider = "International Monetary Fund"
        license_text = "IMF Copyright and Usage terms (statistical data)"
        attribution = (
            "Source: International Monetary Fund, Primary Commodity Price System (PCPS), "
            "https://data.imf.org/"
        )
    return {
        "key": "gold-price",
        "name": name,
        "provider": provider,
        "license": license_text,
        "attribution": attribution,
        "coverage_start": months[0].isoformat(),
        "coverage_end": month_end(months[-1]).isoformat(),
        "as_of": result.as_of.isoformat(),
        "refresh_cadence": "monthly",
    }


def payload(result: PriceResult) -> dict[str, Any]:
    missing = gaps(result.prices)
    return {
        "$generated": generated(
            "gold-price", warnings=result.warnings, missing_months=[f"{m:%Y-%m}" for m in missing]
        ),
        "datasets": [dataset(result)],
        "sources": [result.source.seed()],
        "prices": table(
            ["month", "usd_per_troy_oz", "source"],
            [
                [m.isoformat(), result.prices[m], result.source.source_key]
                for m in sorted(result.prices)
            ],
        ),
    }


def configure(parser: argparse.ArgumentParser) -> None:
    parser.add_argument("--refresh", action="store_true", help="re-fetch instead of the cache")


def run(args: argparse.Namespace) -> int:
    result = build(refresh=bool(args.refresh))
    write_seed(PRICE_SEED, payload(result))
    months = sorted(result.prices)
    print(
        f"Gold price: {len(months)} months from {months[0]:%Y-%m} to {months[-1]:%Y-%m} "
        f"({result.source.title}); latest {result.prices[months[-1]]:,.2f} US$/troy oz."
    )
    for warning in result.warnings:
        print(f"  {warning}")
    return 0
